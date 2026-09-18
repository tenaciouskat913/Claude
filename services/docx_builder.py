import io
import re

from docx import Document
from docx.shared import Pt, RGBColor

import config

AUTHOR_EMAIL = "kat@explicitmathematicsprogram.com"


def slugify(text: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", text).strip("_")
    return slug[:60] or "fact_sheet"


def build_docx(factsheet: dict) -> io.BytesIO:
    doc = Document()

    title = factsheet.get("title", "EMP Policy Fact Sheet")
    doc.core_properties.title = title
    doc.core_properties.author = AUTHOR_EMAIL

    doc.add_heading(title, level=0)

    subtitle_bits = [b for b in [factsheet.get("jurisdiction"), factsheet.get("policy_reference")] if b]
    if subtitle_bits:
        p = doc.add_paragraph(" — ".join(subtitle_bits))
        p.runs[0].italic = True
        p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_heading("Policy Summary", level=2)
    doc.add_paragraph(factsheet.get("policy_summary", ""))

    doc.add_heading("How EMP Supports This Policy", level=2)
    for point in factsheet.get("alignment_points", []):
        para = doc.add_paragraph(style="List Bullet")
        req_run = para.add_run(point.get("policy_requirement", ""))
        req_run.bold = True
        para.add_run(" → " + point.get("emp_feature", "") + ". ")
        para.add_run(point.get("explanation", ""))

    doc.add_heading("Evidence & Outcomes", level=2)
    doc.add_paragraph(factsheet.get("evidence_section", ""))

    doc.add_heading("Scope & Considerations", level=2)
    doc.add_paragraph(factsheet.get("scope_considerations", ""))

    doc.add_paragraph()
    disclaimer_para = doc.add_paragraph()
    disclaimer_run = disclaimer_para.add_run(config.DISCLAIMER_TEXT)
    disclaimer_run.italic = True
    disclaimer_run.font.size = Pt(9)
    disclaimer_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
